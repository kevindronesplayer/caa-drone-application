# ═══════════════════════════════════════════════════════════════
# 民航局無人機申請資料填寫系統 — 後端
#   /api/caa-zones           民航局限制區（UAV_fs）代理＋快取，失敗改用 RCR 離線快照
#   /api/launch-candidates   Overpass 查詢範圍內的公園／戶外停車場
#   /api/reverse             Nominatim 反向地理編碼（空域所在縣市鄉鎮）
#   /api/ai-expand           Claude 擴寫作業概述（串流）
#   /api/export-docx         案名＋作業概述輸出 Word
# ═══════════════════════════════════════════════════════════════

import io
import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
import urllib3
from flask import Flask, Response, jsonify, request, send_file, send_from_directory

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / 'static'


def _load_dotenv():
    """讀取同資料夾的 .env（只處理 KEY=VALUE），讓使用者不必設定系統環境變數"""
    env_path = BASE_DIR / '.env'
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue
        key, value = line.split('=', 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


_load_dotenv()
os.environ.setdefault('FLASK_SKIP_DOTENV', '1')  # .env 已自行讀取，不需 Flask 再提示安裝 python-dotenv

app = Flask(__name__, static_folder=None)
app.json.ensure_ascii = False

BROWSER_UA = ('Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 '
              '(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36')
APP_UA = 'CAA-Drone-Application-Helper/1.0 (local desktop tool)'


# ── 民航局限制區 ──────────────────────────────────────────────
# 與 dronemap 筆記本相同的公開 REST 端點（dronegis.caa.gov.tw 公眾版地圖所用）
CAA_UAV_QUERY_URL = 'https://dronegis.caa.gov.tw/server/rest/services/Hosted/UAV_fs/FeatureServer/3/query'
CAA_CACHE_TTL = 3600
_caa_cache = {'time': 0, 'data': None}


def _fetch_caa_zones(page_size=2000, timeout=25):
    headers = {
        'User-Agent': BROWSER_UA,
        'Referer': 'https://dronegis.caa.gov.tw/portal/apps/webappviewer/index.html?id=807bd21438ba4208b4a7e28569fe41aa',
        'Accept': 'application/json, text/plain, */*',
    }
    features, offset, verify = [], 0, True
    while True:
        params = {
            'where': '1=1',
            'outFields': '空域名稱,空域類別名稱,空域顏色,有效日期起,有效日期迄,條件,主管機關名稱',
            'returnGeometry': 'true',
            'outSR': '4326',
            'f': 'geojson',
            'maxAllowableOffset': '0.0003',
            'geometryPrecision': '5',
            'resultOffset': str(offset),
            'resultRecordCount': str(page_size),
        }
        try:
            resp = requests.get(CAA_UAV_QUERY_URL, params=params, headers=headers,
                                timeout=timeout, verify=verify)
        except requests.exceptions.SSLError:
            if not verify:
                raise
            # 民航局憑證鏈的 TWCA 根憑證缺 SKI 欄位，部分 OpenSSL 會擋；公開圖資改不驗證重試
            verify = False
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
            continue
        resp.raise_for_status()
        page = resp.json()
        feats = page.get('features', [])
        features.extend(feats)
        if len(feats) < page_size:
            break
        offset += page_size
    return {'type': 'FeatureCollection', 'features': features}


@app.get('/api/caa-zones')
def caa_zones():
    now = time.time()
    if _caa_cache['data'] and now - _caa_cache['time'] < CAA_CACHE_TTL:
        return jsonify(_caa_cache['data'])
    try:
        data = _fetch_caa_zones()
        payload = {'source': 'live', 'fetched_at': int(now), **data}
        _caa_cache.update(time=now, data=payload)
        return jsonify(payload)
    except Exception as e:  # 離線或被擋 → 改用內建 RCR 快照
        fallback = json.loads((STATIC_DIR / 'rcr_fallback.geojson').read_text(encoding='utf-8'))
        return jsonify({'source': 'fallback', 'error': str(e), **fallback})


# ── 起飛點候選（公園／戶外停車場）──────────────────────────────
OVERPASS_URLS = [
    'https://overpass-api.de/api/interpreter',
    'https://overpass.kumi.systems/api/interpreter',
    'https://maps.mail.ru/osm/tools/overpass/api/interpreter',
]
_overpass_cache = {}
INDOOR_PARKING_WORDS = ('地下', '立體', '室內', '大樓', 'B1', 'B2')


def _classify_candidate(tags):
    if tags.get('amenity') == 'parking':
        return 'parking'
    return 'park'


@app.post('/api/launch-candidates')
def launch_candidates():
    body = request.get_json(force=True)
    s, w, n, e = (float(body[k]) for k in ('south', 'west', 'north', 'east'))
    key = tuple(round(v, 4) for v in (s, w, n, e))
    if key in _overpass_cache:
        return jsonify(_overpass_cache[key])

    bbox = f'{s},{w},{n},{e}'
    query = f"""
[out:json][timeout:25];
(
  nwr["leisure"~"^(park|recreation_ground|garden)$"]["access"!~"^(private|no)$"]({bbox});
  nwr["amenity"="parking"]["parking"!~"^(underground|multi-storey|rooftop)$"]["location"!~"underground"]["access"!~"^(private|no)$"]({bbox});
);
out center tags 500;
"""
    # Overpass 各鏡像站速度時好時壞：同時查詢，取最先成功的結果
    def query_mirror(url):
        resp = requests.post(url, data={'data': query}, headers={'User-Agent': APP_UA}, timeout=30)
        resp.raise_for_status()
        return resp.json().get('elements', [])

    elements, errors = None, []
    pool = ThreadPoolExecutor(max_workers=len(OVERPASS_URLS))
    futures = [pool.submit(query_mirror, url) for url in OVERPASS_URLS]
    for fut in as_completed(futures):
        try:
            elements = fut.result()
            break
        except Exception as err:
            errors.append(str(err))
    pool.shutdown(wait=False, cancel_futures=True)
    if elements is None:
        return jsonify({'error': f'Overpass 查詢失敗：{errors[-1] if errors else "未知錯誤"}', 'candidates': []}), 502

    candidates = []
    for el in elements:
        lat = el.get('lat') or (el.get('center') or {}).get('lat')
        lon = el.get('lon') or (el.get('center') or {}).get('lon')
        if lat is None or lon is None:
            continue
        tags = el.get('tags', {})
        kind = _classify_candidate(tags)
        name = tags.get('name:zh') or tags.get('name') or ''
        # 只要戶外停車場：有些地下／立體停車場沒標 parking 類型，只寫在名稱裡
        if kind == 'parking' and any(k in name for k in INDOOR_PARKING_WORDS):
            continue
        if not name:
            name = '戶外停車場' if kind == 'parking' else '公園綠地'
        candidates.append({
            'id': f"{el['type']}/{el['id']}",
            'lat': lat, 'lng': lon, 'kind': kind,
            'name': name, 'named': bool(tags.get('name')),
        })
    result = {'candidates': candidates}
    _overpass_cache[key] = result
    return jsonify(result)


# ── 反向地理編碼 ──────────────────────────────────────────────
_reverse_cache = {}


@app.get('/api/reverse')
def reverse_geocode():
    lat, lng = float(request.args['lat']), float(request.args['lng'])
    key = (round(lat, 4), round(lng, 4))
    if key in _reverse_cache:
        return jsonify(_reverse_cache[key])
    try:
        resp = requests.get('https://nominatim.openstreetmap.org/reverse', params={
            'lat': lat, 'lon': lng, 'format': 'jsonv2', 'zoom': 16, 'accept-language': 'zh-TW',
        }, headers={'User-Agent': APP_UA}, timeout=15)
        resp.raise_for_status()
        addr = resp.json().get('address', {})
    except Exception as err:
        return jsonify({'error': str(err), 'area': '', 'place': ''}), 502
    city = addr.get('city') or addr.get('county') or addr.get('state') or ''
    town = addr.get('town') or addr.get('suburb') or addr.get('city_district') or addr.get('village') or ''
    if town == city:
        town = ''
    place = addr.get('amenity') or addr.get('leisure') or addr.get('road') or ''
    result = {'area': f'{city}{town}', 'place': place}
    _reverse_cache[key] = result
    return jsonify(result)


@app.get('/api/search')
def search_place():
    q = (request.args.get('q') or '').strip()
    if not q:
        return jsonify({'results': []})
    try:
        resp = requests.get('https://nominatim.openstreetmap.org/search', params={
            'q': q, 'format': 'jsonv2', 'limit': 5, 'countrycodes': 'tw', 'accept-language': 'zh-TW',
        }, headers={'User-Agent': APP_UA}, timeout=15)
        resp.raise_for_status()
        items = resp.json()
    except Exception as err:
        return jsonify({'error': str(err), 'results': []}), 502
    results = []
    for it in items:
        bbox = it.get('boundingbox')
        results.append({
            'name': it.get('display_name', ''),
            'lat': float(it['lat']), 'lng': float(it['lon']),
            'bbox': [float(v) for v in bbox] if bbox else None,  # [south, north, west, east]
        })
    return jsonify({'results': results})


# ── AI 擴寫作業概述 ───────────────────────────────────────────
AI_MODEL = os.environ.get('CLAUDE_MODEL', 'claude-opus-5-5')

AI_SYSTEM_PROMPT = """你是協助台灣無人機業者撰寫「交通部民用航空局 遙控無人機活動申請」文件的專業文書助理。
使用者會提供案名、作業大致內容，以及系統自動整理的空域資料。請把大致內容擴寫成一份正式、完整、可直接貼進申請書的「作業概述」。

撰寫要求：
1. 使用台灣正體中文與公文書常用語氣，條理分明。
2. 以「一、二、三…」分段，建議包含：作業目的、作業地點及空域範圍、作業期間及時段、飛航高度、作業方式及飛行規劃、人員配置與安全管理措施、緊急應變措施。可依內容增減，但不要灌水。
3. 只能使用使用者與空域資料中的事實。日期、時段、機型、操作人員證號、保險等未提供的資訊，一律寫成【待確認：項目】佔位，絕對不要自行編造。
4. 不要寫「預計起飛地點」段落（系統會另外插入）；若已提供起飛點資料，可在作業方式中簡要提及「詳見預計起飛地點」。
5. 輸出純文字，不要使用 Markdown 符號（#、*、**、- 項目符號、表格），因為內容會直接輸出成 Word。
6. 直接輸出作業概述本文，不要前言或結語說明。"""


def _anthropic_client():
    import anthropic
    return anthropic.Anthropic()


@app.post('/api/ai-expand')
def ai_expand():
    import anthropic

    body = request.get_json(force=True)
    case_name = (body.get('case_name') or '').strip()
    draft = (body.get('draft') or '').strip()
    airspace_text = (body.get('airspace_summary') or '').strip()
    if not draft:
        return jsonify({'error': '請先輸入作業大致內容'}), 400

    user_msg = (f'案名：{case_name or "【待確認：案名】"}\n\n'
                f'作業大致內容：\n{draft}\n\n'
                f'空域資料（系統自動整理）：\n{airspace_text or "（無）"}')

    try:
        client = _anthropic_client()
        manager = client.beta.messages.stream(
            model=AI_MODEL,
            max_tokens=16000,
            system=AI_SYSTEM_PROMPT,
            messages=[{'role': 'user', 'content': user_msg}],
            output_config={'effort': 'medium'},
            betas=['server-side-fallback-2026-07-01'],
            fallbacks='default',
        )
        stream = manager.__enter__()
    except anthropic.AuthenticationError:
        return jsonify({'error': 'Claude API 金鑰無效，請檢查 .env 中的 ANTHROPIC_API_KEY'}), 401
    except anthropic.APIConnectionError as err:
        return jsonify({'error': f'無法連線 Claude API：{err}'}), 502
    except anthropic.APIStatusError as err:
        return jsonify({'error': f'Claude API 錯誤（{err.status_code}）：{err.message}'}), 502
    except Exception as err:  # 多半是沒有設定金鑰
        return jsonify({'error': f'尚未設定 Claude API 金鑰（{err}）。請在專案資料夾的 .env 填入 ANTHROPIC_API_KEY'}), 500

    def generate():
        try:
            for text in stream.text_stream:
                yield text
            final = stream.get_final_message()
            if final.stop_reason == 'refusal':
                yield '\n\n【AI 拒絕產生此內容，請修改大致內容後再試】'
            elif final.stop_reason == 'max_tokens':
                yield '\n\n【內容過長被截斷】'
        except anthropic.APIError as err:
            yield f'\n\n【AI 產生中斷：{err}】'
        finally:
            manager.__exit__(None, None, None)

    return Response(generate(), mimetype='text/plain; charset=utf-8',
                    headers={'X-Accel-Buffering': 'no', 'Cache-Control': 'no-cache'})


# ── Word 輸出 ────────────────────────────────────────────────
def _set_run_font(run, size_pt, bold=False, font='標楷體'):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    for attr in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rfonts.set(qn(attr), font)


@app.post('/api/export-docx')
def export_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    body = request.get_json(force=True)
    case_name = (body.get('case_name') or '').strip() or '未命名案件'
    overview = (body.get('overview') or '').replace('\r\n', '\n').strip()

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title.add_run(case_name), 18, bold=True)
    title.paragraph_format.space_after = Pt(18)

    label = doc.add_paragraph()
    _set_run_font(label.add_run('案名：'), 14, bold=True)
    _set_run_font(label.add_run(case_name), 14)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(12)
    _set_run_font(heading.add_run('作業概述：'), 14, bold=True)

    # 「一、」「【預計起飛地點】」這類段落標題加粗，其餘照原文換行輸出
    section_re = re.compile(r'^([一二三四五六七八九十]+、|【)')
    for line in overview.split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        _set_run_font(p.add_run(line), 12, bold=bool(section_re.match(line.strip())))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    safe_name = re.sub(r'[\\/:*?"<>|]', '_', case_name)
    return send_file(buf, as_attachment=True, download_name=f'{safe_name}_作業概述.docx',
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


@app.get('/api/status')
def status():
    has_key = bool(os.environ.get('ANTHROPIC_API_KEY') or os.environ.get('ANTHROPIC_AUTH_TOKEN'))
    return jsonify({'ai_ready': has_key, 'model': AI_MODEL})


# ── 靜態檔 ───────────────────────────────────────────────────
@app.get('/')
def index():
    return send_from_directory(STATIC_DIR, 'index.html')


@app.get('/<path:path>')
def static_files(path):
    return send_from_directory(STATIC_DIR, path)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 8765))
    print(f'🚁 民航局無人機申請填寫系統：http://127.0.0.1:{port}')
    app.run(host='127.0.0.1', port=port, debug=False, threaded=True)
